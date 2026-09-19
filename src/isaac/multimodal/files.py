from __future__ import annotations

import base64
import io
import tempfile
import zipfile
from itertools import islice
from pathlib import Path

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TOTAL_BYTES = 8 * 1024 * 1024
MAX_ATTACHMENTS = 4
MAX_TEXT_CHARS = 60_000
MAX_IMAGE_PIXELS = 16_000_000
TEXT_SUFFIXES = frozenset(
    [
        ".txt",
        ".md",
        ".markdown",
        ".py",
        ".js",
        ".ts",
        ".tsx",
        ".jsx",
        ".json",
        ".jsonl",
        ".yaml",
        ".yml",
        ".toml",
        ".ini",
        ".cfg",
        ".csv",
        ".tsv",
        ".log",
        ".html",
        ".htm",
        ".css",
        ".xml",
        ".rs",
        ".go",
        ".java",
        ".c",
        ".h",
        ".cpp",
        ".sh",
        ".ps1",
        ".sql",
        ".r",
        ".gitignore",
        ".dockerfile",
    ]
)
IMAGE_SUFFIXES = frozenset({".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"})
AUDIO_SUFFIXES = frozenset({".mp3", ".wav", ".m4a", ".ogg", ".flac"})


class AttachmentError(ValueError):
    pass


def describe_file(path: str | Path) -> dict:
    p = Path(path)
    try:
        return {
            "path": str(p),
            "exists": p.is_file(),
            "name": p.name,
            "suffix": p.suffix.lower(),
            "size_bytes": p.stat().st_size,
        }
    except OSError:
        return {"path": str(p), "exists": False}


def read_any_file(path: str | Path, max_chars: int = MAX_TEXT_CHARS) -> dict:
    meta = describe_file(path)
    try:
        p = Path(path).expanduser()
        if not p.is_file():
            raise AttachmentError("File not found or not a regular file")
        if p.stat().st_size > MAX_FILE_BYTES:
            raise AttachmentError(f"File exceeds {MAX_FILE_BYTES} bytes")
        with p.open("rb") as stream:
            raw = stream.read(MAX_FILE_BYTES + 1)
        result = parse_file(p.name, raw, max_chars=max_chars)
        result["meta"] = meta
        return result
    except (OSError, ValueError) as exc:
        return {
            "ok": False,
            "error": str(exc),
            "text": f"[Attachment error: {exc}]",
            "images": [],
            "attachments": [],
            "meta": meta,
        }


def parse_file(name: str, raw: bytes, *, max_chars: int = MAX_TEXT_CHARS) -> dict:
    if not isinstance(name, str) or not name or len(name) > 255:
        raise AttachmentError("Invalid attachment filename")
    name = name.replace("\\", "/").rsplit("/", 1)[-1]
    if not name or any(ord(c) < 32 for c in name):
        raise AttachmentError("Invalid attachment filename")
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise AttachmentError(f"Attachment must contain 1–{MAX_FILE_BYTES} bytes")
    suffix = Path(name).suffix.lower()
    max_chars = max(1, min(max_chars, MAX_TEXT_CHARS))
    images: list[str] = []
    try:
        if suffix in IMAGE_SUFFIXES:
            images = [_image_url(raw)]
            text = f"[Image: {name}]"
        elif suffix in AUDIO_SUFFIXES:
            text = _read_audio(raw, suffix)
        elif suffix == ".pdf":
            text = _read_pdf(raw, max_chars)
        elif suffix in {".docx", ".xlsx", ".pptx"}:
            _check_archive(raw)
            text = _read_office(raw, suffix, max_chars)
        elif suffix in TEXT_SUFFIXES or name.lower() in {"dockerfile", ".gitignore"}:
            text = raw.decode("utf-8-sig", errors="strict")
            if "\0" in text:
                raise AttachmentError("Binary content is not supported as text")
        else:
            raise AttachmentError(f"Unsupported file format: {suffix or '(no extension)'} (binary)")
    except ImportError as exc:
        raise AttachmentError(f"Missing parser dependency: {exc}") from exc
    except AttachmentError:
        raise
    except Exception as exc:
        raise AttachmentError(f"Could not parse {name}: {exc}") from exc
    if not text.strip():
        raise AttachmentError(
            "No text extracted; scanned documents require OCR (not supported here)"
        )
    clipped = len(text) > max_chars
    text = f"[File: {name}]\n{text[:max_chars]}" + ("\n[truncated]" if clipped else "")
    blocks = [{"type": "text", "text": text}]
    blocks.extend({"type": "image_url", "image_url": {"url": url}} for url in images)
    return {
        "ok": True,
        "text": text,
        "images": images,
        "attachments": blocks,
        "meta": {"name": name, "size_bytes": len(raw), "suffix": suffix},
    }


def parse_uploads(uploads: object) -> list[dict]:
    if not isinstance(uploads, list) or len(uploads) > MAX_ATTACHMENTS:
        raise AttachmentError(f"Provide at most {MAX_ATTACHMENTS} attachments")
    blocks: list[dict] = []
    total = 0
    for upload in uploads:
        if not isinstance(upload, dict):
            raise AttachmentError("Invalid attachment")
        data = upload.get("data")
        if not isinstance(data, str) or len(data) > 4 * ((MAX_FILE_BYTES + 2) // 3):
            raise AttachmentError("Attachment base64 exceeds size limit")
        try:
            raw = base64.b64decode(data, validate=True)
        except ValueError as exc:
            raise AttachmentError("Invalid attachment base64") from exc
        total += len(raw)
        if total > MAX_TOTAL_BYTES:
            raise AttachmentError("Total attachment size exceeds limit")
        blocks.extend(parse_file(upload.get("name"), raw)["attachments"])
    return blocks


def _image_url(raw: bytes) -> str:
    from PIL import Image

    with Image.open(io.BytesIO(raw)) as image:
        if image.width * image.height > MAX_IMAGE_PIXELS:
            raise AttachmentError("Image exceeds pixel limit")
        if image.format not in {"PNG", "JPEG", "WEBP", "GIF", "BMP"}:
            raise AttachmentError("Unsupported image encoding")
        if getattr(image, "n_frames", 1) != 1:
            raise AttachmentError("Animated images are not supported")
        image.verify()
    with Image.open(io.BytesIO(raw)) as image:
        out = io.BytesIO()
        image.convert("RGB").save(out, format="PNG")
    normalized = out.getvalue()
    if len(normalized) > MAX_FILE_BYTES:
        raise AttachmentError("Decoded image exceeds size limit")
    return "data:image/png;base64," + base64.b64encode(normalized).decode("ascii")


def _read_pdf(raw: bytes, limit: int) -> str:
    try:
        import fitz
    except ImportError as exc:
        raise AttachmentError("PDF parser unavailable; install pymupdf") from exc
    with fitz.open(stream=raw, filetype="pdf") as doc:
        if doc.needs_pass:
            raise AttachmentError("Password-protected PDFs are not supported")
        if doc.page_count > 200:
            raise AttachmentError("PDF exceeds 200-page limit")
        return _collect((page.get_text() for page in doc), limit)


def _check_archive(raw: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 2048 or sum(e.file_size for e in entries) > 32 * 1024 * 1024:
            raise AttachmentError("Document exceeds expanded archive limit")


def _collect(parts, limit: int) -> str:
    out: list[str] = []
    remaining = limit + 1
    for part in parts:
        chunk = str(part)[:remaining]
        out.append(chunk)
        remaining -= len(chunk) + 1
        if remaining <= 0:
            break
    return "\n".join(out)


def _read_office(raw: bytes, suffix: str, limit: int) -> str:
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise AttachmentError("DOCX parser unavailable; install python-docx") from exc
        return _collect((p.text for p in Document(io.BytesIO(raw)).paragraphs), limit)
    if suffix == ".pptx":
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise AttachmentError("PPTX parser unavailable; install python-pptx") from exc
        return _collect(
            (
                shape.text
                for slide in Presentation(io.BytesIO(raw)).slides
                for shape in slide.shapes
                if shape.has_text_frame
            ),
            limit,
        )
    try:
        import openpyxl
    except ImportError as exc:
        raise AttachmentError("XLSX parser unavailable; install openpyxl") from exc
    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        return _collect(
            (
                ", ".join(str(v) if v is not None else "" for v in row)
                for sheet in workbook.worksheets[:5]
                for row in islice(sheet.iter_rows(max_col=100, values_only=True), 500)
            ),
            limit,
        )
    finally:
        workbook.close()


def _read_audio(raw: bytes, suffix: str) -> str:
    try:
        import av
    except ImportError as exc:
        raise AttachmentError(
            "Audio parser unavailable; install faster-whisper (includes av)"
        ) from exc
    with av.open(io.BytesIO(raw)) as container:
        if not container.streams.audio or container.duration is None:
            raise AttachmentError("Audio duration could not be verified")
        if container.duration / av.time_base > 300:
            raise AttachmentError("Audio exceeds five-minute limit")
    from isaac.multimodal.voice.stt import get_stt

    with tempfile.TemporaryDirectory(prefix="isaac-attachment-") as directory:
        path = Path(directory) / f"audio{suffix}"
        path.write_bytes(raw)
        return get_stt().transcribe(str(path))
